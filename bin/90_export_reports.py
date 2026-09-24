#!/usr/bin/env python3
"""bin/90: export human-readable reports (Excel, Word, text) into reports/.

reports/ is gitignored; sync_onedrive.sh copies it to the OneDrive project folder.
Currently exports the genome inventory and provenance audit. Oligo-set tables are added
here once bin/12 produces oligos.

Needs: openpyxl (Excel), python-docx (Word). Both are in the base environment.
"""
import collections
import csv
import datetime
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
INV = ROOT / "data" / "genome_inventory"
OUT = ROOT / "reports"

TIER_ORDER = {"INCLUDE": 0, "QUALITY_REVIEW": 1, "REVIEW": 2, "EXCLUDE": 3}
TIER_LABEL = {"INCLUDE": "primary", "QUALITY_REVIEW": "quality-review", "REVIEW": "UNRESOLVED", "EXCLUDE": "excluded"}
TIER_FILL = {"primary": "E2F0D9", "quality-review": "FFF2CC", "UNRESOLVED": "F8CBAD", "excluded": "F4CCCC"}


def read_tsv(path):
    return list(csv.DictReader(open(path), delimiter="\t")) if path.exists() else []


def load():
    meta = {r["accession"]: r for r in read_tsv(INV / "genome_metadata.tsv")}
    audit = {r["accession"]: r for r in read_tsv(INV / "provenance_audit.tsv")}
    files = {r["accession"]: r for r in read_tsv(INV / "genome_files.tsv")}
    if not meta or not audit:
        sys.exit("run bin/01 and bin/02 first")
    rows = []
    for acc, m in meta.items():
        a, f = audit[acc], files.get(acc, {})
        tier = TIER_LABEL[a["final_verdict"]]
        reason = a["override_reason"] or a["hard_flags"] or a["strong_flags"]
        rows.append({
            "Accession": acc, "Strain": m["strain"], "Serovar": m["serovar"], "Tier": tier,
            "BioProject": m["bioproject"], "BioProject title": m["bioproject_titles"],
            "Collection date": m["collection_date"], "Location": m["geo_loc_name"], "Host": m["host"],
            "Isolation source": m["isolation_source"], "Total length (nt)": int(m["total_length"] or 0),
            "Chromosome length (nt)": int(f["longest_seq"]) if f.get("longest_seq") else "",
            "Plasmid in assembly": ("yes" if f.get("small_replicons", "0") not in ("", "0") else "no") if f else "",
            "GC %": float(m["gc_percent"]) if m["gc_percent"] else "",
            "CheckM completeness": float(m["checkm_completeness"]) if m["checkm_completeness"] else "",
            "CheckM contamination": float(m["checkm_contamination"]) if m["checkm_contamination"] else "",
            "Decision reason / flags": reason, "Notes": a["weak_notes"],
            "_order": TIER_ORDER[a["final_verdict"]],
        })
    rows.sort(key=lambda r: (r["_order"], r["Strain"].lower()))
    return rows, meta, audit


def write_xlsx(rows, meta, path):
    wb = Workbook()
    cols = [c for c in rows[0] if not c.startswith("_")]
    ws = wb.active
    ws.title = "Genomes"
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    tier_col = cols.index("Tier") + 1
    for row in ws.iter_rows(min_row=2):
        fill = PatternFill("solid", fgColor=TIER_FILL[row[tier_col - 1].value])
        row[tier_col - 1].fill = fill
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"Accession": 17, "Strain": 22, "Serovar": 9, "Tier": 14, "BioProject": 14, "BioProject title": 44,
              "Decision reason / flags": 70, "Notes": 40, "Isolation source": 22, "Location": 20}
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 13)

    bp = collections.defaultdict(lambda: collections.Counter())
    info = {}
    for r in rows:
        bp[r["BioProject"]][r["Tier"]] += 1
        info.setdefault(r["BioProject"], (r["BioProject title"], meta[r["Accession"]]["submitter"]))
    ws2 = wb.create_sheet("BioProjects")
    ws2.append(["BioProject", "Genomes", "primary", "quality-review", "excluded", "UNRESOLVED", "Submitter", "Title"])
    for k, c in sorted(bp.items(), key=lambda kv: -sum(kv[1].values())):
        ws2.append([k, sum(c.values()), c["primary"], c["quality-review"], c["excluded"], c["UNRESOLVED"], info[k][1], info[k][0]])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for i, w in enumerate([14, 9, 9, 14, 9, 12, 32, 90], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    counts = collections.Counter(r["Tier"] for r in rows)
    ws3 = wb.create_sheet("Summary")
    ws3.append(["C. trachomatis RefSeq complete genome inventory"])
    ws3["A1"].font = Font(bold=True, size=13)
    ws3.append([f"Generated {datetime.date.today()} by bin/90_export_reports.py"])
    ws3.append([])
    ws3.append(["Tier", "Genomes", "Meaning"])
    meaning = {
        "primary": "Used for consensus building and the inclusivity claim",
        "quality-review": "Natural isolates with suspicious quality metrics: tested for inclusivity, not used for consensus",
        "excluded": "Experimental, engineered or laboratory-selected strains",
        "UNRESOLVED": "Flagged, awaiting a decision in config/provenance_overrides.tsv",
    }
    for t in ("primary", "quality-review", "excluded", "UNRESOLVED"):
        ws3.append([t, counts[t], meaning[t]])
    ws3.append(["Total", sum(counts.values())])
    for c in ws3[4]:
        c.font = Font(bold=True)
    ws3.column_dimensions["A"].width = 18
    ws3.column_dimensions["C"].width = 90
    wb.move_sheet("Summary", offset=-2)
    wb.save(path)


def add_table(doc, header, body, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for row in body:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)
    return t


def write_docx(rows, path):
    counts = collections.Counter(r["Tier"] for r in rows)
    doc = Document()
    doc.add_heading("C. trachomatis genome inventory and provenance audit", level=0)
    doc.add_paragraph(f"Generated {datetime.date.today()} from the restart pipeline (bin/01, bin/02). "
                      "Full per-genome table: genome_inventory.xlsx.")
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(
        f"{sum(counts.values())} RefSeq complete C. trachomatis genomes were retrieved from NCBI and audited for provenance "
        f"before any analysis. {counts['primary']} are in the primary tier, {counts['quality-review']} in the quality-review tier "
        f"and {counts['excluded']} were excluded" + (f"; {counts['UNRESOLVED']} are unresolved." if counts["UNRESOLVED"] else "."))
    doc.add_heading("Tiers", level=1)
    doc.add_paragraph("Primary: used for consensus building and the inclusivity claim.", style="List Bullet")
    doc.add_paragraph("Quality-review: natural isolates with suspicious quality metrics. Not used to build consensus; every oligo "
                      "is tested against them and the results are reported separately.", style="List Bullet")
    doc.add_paragraph("Excluded: experimental, engineered or laboratory-selected strains, and anything shown not to be "
                      "C. trachomatis.", style="List Bullet")
    for tier, title in (("excluded", "Excluded genomes"), ("quality-review", "Quality-review genomes"), ("UNRESOLVED", "Unresolved genomes")):
        sel = [r for r in rows if r["Tier"] == tier]
        if not sel:
            continue
        doc.add_heading(f"{title} ({len(sel)})", level=1)
        add_table(doc, ["Accession", "Strain", "BioProject", "Reason"],
                  [[r["Accession"], r["Strain"], r["BioProject"], r["Decision reason / flags"][:220]] for r in sel])
    doc.add_heading("Method notes", level=1)
    for text in (
        "Every genome is traced to its BioProject and BioSample metadata before use. 'Complete genome' describes assembly "
        "contiguity, not whether an isolate is natural.",
        "Automatic rules flag engineering or selection markers in names and descriptions, experiment-style BioProject titles, "
        "ANI and CheckM anomalies, and chromosome length more than 8 kb from the median. Manual decisions are recorded "
        "with reasons in config/provenance_overrides.tsv.",
        "Known-bad BioProjects (PRJNA558398 lab-engineered recombinants; PRJEB35640 metagenome-assembled artifacts) are "
        "excluded by rule; neither appears in the RefSeq complete set.",
    ):
        doc.add_paragraph(text, style="List Bullet")
    doc.save(path)


def main():
    OUT.mkdir(exist_ok=True)
    rows, meta, audit = load()
    write_xlsx(rows, meta, OUT / "genome_inventory.xlsx")
    write_docx(rows, OUT / "genome_provenance_report.docx")
    log = INV / "provenance_audit_log.txt"
    if log.exists():
        shutil.copy(log, OUT / "provenance_audit_log.txt")
    print("wrote:")
    for p in sorted(OUT.iterdir()):
        print(f"  reports/{p.name} ({p.stat().st_size:,} bytes)")
    print(dict(collections.Counter(r["Tier"] for r in rows)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
